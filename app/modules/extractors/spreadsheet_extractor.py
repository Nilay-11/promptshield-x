"""
PromptShield X - Spreadsheet Extractor
Parses Excel (.xlsx/.xls) and CSV files, detects DDE/Formula injection, and hidden sheet payloads.
Uses Pandas for tabular data extraction and OpenPyXL for structural metadata extraction.
"""

import io
import openpyxl
import pandas as pd
from typing import Optional
from app.modules.extractors import BaseExtractor, ExtractedSegment, ExtractionResult, SourceType, ThreatCategory


class SpreadsheetExtractor(BaseExtractor):
    DDE_PREFIXES = ("=", "@", "+", "-", "|", "\t", "\r")

    def extract(self, data: bytes, filename: Optional[str] = None, **kwargs) -> ExtractionResult:
        result = ExtractionResult(source_type=SourceType.SPREADSHEET, filename=filename)
        is_csv = filename.lower().endswith(".csv") if filename else False

        if is_csv:
            self._extract_csv(data, result)
        else:
            self._extract_excel(data, result, filename)

        result.raw_character_count = sum(len(s.content) for s in result.segments)
        result.anomalies_detected = sum(1 for s in result.segments if s.is_hidden or s.threat_indicators)
        return result

    def _get_threats(self, text: str) -> tuple[list[ThreatCategory], float]:
        threats = []
        penalty = 0.0
        if text.startswith(self.DDE_PREFIXES):
            threats.append(ThreatCategory.FORMULA_INJECTION)
            penalty += 35.0
        return threats, penalty

    def _extract_csv(self, data: bytes, result: ExtractionResult):
        try:
            # Use Pandas for tabular extraction
            df = pd.read_csv(io.BytesIO(data))
            for row_idx, row in df.iterrows():
                for col_idx, value in enumerate(row):
                    if pd.notna(value) and str(value).strip():
                        val_str = str(value).strip()
                        threats, penalty = self._get_threats(val_str)
                        
                        # Column can be letter or index, here df.columns[col_idx] is the column name
                        result.segments.append(
                            ExtractedSegment(
                                content=val_str,
                                source_type=SourceType.SPREADSHEET,
                                location={
                                    "row": row_idx + 2, # +2 assuming headers take row 1
                                    "column": str(df.columns[col_idx])
                                },
                                is_hidden=False,
                                threat_indicators=threats,
                                confidence_penalty=penalty
                            )
                        )
        except Exception as e:
            result.extraction_warnings.append(f"CSV Pandas parsing error: {str(e)}")

    def _extract_excel(self, data: bytes, result: ExtractionResult, filename: Optional[str]):
        # 1. Structural inspection via OpenPyXL
        try:
            wb = openpyxl.load_workbook(filename=io.BytesIO(data), data_only=False, read_only=False)
            
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                is_sheet_hidden = ws.sheet_state != "visible"
                
                for row_idx, row in enumerate(ws.iter_rows()):
                    row_hidden = ws.row_dimensions[row_idx + 1].hidden
                    
                    for cell in row:
                        cell_coord = cell.coordinate
                        col_hidden = (cell.column_letter in ws.column_dimensions and ws.column_dimensions[cell.column_letter].hidden)
                        
                        is_hidden = is_sheet_hidden or row_hidden or col_hidden
                        
                        # Check Camouflaged Font
                        if cell.fill and cell.fill.start_color and cell.font and cell.font.color:
                            if cell.fill.start_color.rgb == cell.font.color.rgb:
                                is_hidden = True

                        # Extract Comments
                        if cell.comment:
                            c_text = cell.comment.text.strip()
                            if c_text:
                                threats, penalty = self._get_threats(c_text)
                                if is_hidden:
                                    threats.append(ThreatCategory.HIDDEN_TEXT)
                                    penalty += 30.0
                                result.segments.append(
                                    ExtractedSegment(
                                        content=c_text,
                                        source_type=SourceType.SPREADSHEET,
                                        location={"sheet": sheetname, "cell": cell_coord, "element": "comment"},
                                        is_hidden=is_hidden,
                                        threat_indicators=threats,
                                        confidence_penalty=penalty
                                    )
                                )

                        # Extract Hyperlinks
                        if cell.hyperlink and cell.hyperlink.target:
                            h_target = cell.hyperlink.target.strip()
                            threats, penalty = self._get_threats(h_target)
                            if is_hidden:
                                threats.append(ThreatCategory.HIDDEN_TEXT)
                                penalty += 30.0
                            result.segments.append(
                                ExtractedSegment(
                                    content=h_target,
                                    source_type=SourceType.SPREADSHEET,
                                    location={"sheet": sheetname, "cell": cell_coord, "element": "hyperlink"},
                                    is_hidden=is_hidden,
                                    threat_indicators=threats,
                                    confidence_penalty=penalty
                                )
                            )
                            
                        # Extract Formulas (openpyxl retains formula strings since data_only=False)
                        if cell.data_type == 'f' and cell.value:
                            f_val = str(cell.value).strip()
                            # It's a formula, explicitly mark it
                            threats = [ThreatCategory.FORMULA_INJECTION]
                            penalty = 40.0
                            if is_hidden:
                                threats.append(ThreatCategory.HIDDEN_TEXT)
                                penalty += 30.0
                            result.segments.append(
                                ExtractedSegment(
                                    content=f_val,
                                    source_type=SourceType.SPREADSHEET,
                                    location={"sheet": sheetname, "cell": cell_coord, "element": "formula"},
                                    is_hidden=is_hidden,
                                    threat_indicators=threats,
                                    confidence_penalty=penalty
                                )
                            )
            wb.close()
        except Exception as e:
            result.extraction_warnings.append(f"Excel OpenPyXL structured parsing error: {str(e)}")

        # 2. Tabular inspection via Pandas
        try:
            xls = pd.ExcelFile(io.BytesIO(data))
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name)
                for row_idx, row in df.iterrows():
                    for col_idx, value in enumerate(row):
                        if pd.notna(value) and str(value).strip():
                            val_str = str(value).strip()
                            # We might have formula injection in raw string format
                            threats, penalty = self._get_threats(val_str)
                            
                            col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
                            cell_coord = f"{col_letter}{row_idx + 2}"
                            
                            result.segments.append(
                                ExtractedSegment(
                                    content=val_str,
                                    source_type=SourceType.SPREADSHEET,
                                    location={"sheet": sheet_name, "cell": cell_coord},
                                    is_hidden=False, # Pandas doesn't know about visibility, handled by OpenPyXL
                                    threat_indicators=threats,
                                    confidence_penalty=penalty
                                )
                            )
        except Exception as e:
            result.extraction_warnings.append(f"Excel Pandas extraction error: {str(e)}")